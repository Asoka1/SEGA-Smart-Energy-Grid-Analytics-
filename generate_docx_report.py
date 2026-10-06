"""Generate submission-ready Word document report for Assignment I."""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def main():
    doc = Document()

    # Style definitions
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Smart Energy Grid Analytics (SEGA)\n")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    run_sub = p_title.add_run("Advanced Big Data Analytics — Assignment I Technical Report\n")
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.add_run("Adventist University of Central Africa (AUCA) | October 2026\nTimezone: Africa/Kigali (CAT, UTC+02:00)\n")

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This report documents the design, benchmarking, and analytical findings for the Smart Energy Grid "
        "Analytics (SEGA) system covering 4,000 smart meters over a 28-day historical period in Rwanda. "
        "Using TimescaleDB on PostgreSQL 16 and EMQX, all data ingestion, chunk size comparisons, "
        "compression strategies, continuous aggregations, and day-ahead demand forecasts were executed and rigorously evaluated."
    )

    doc.add_heading("2. Ingestion Pipeline & Verification", level=1)
    doc.add_paragraph(
        "The environment was confirmed with EMQX v6.3.1 and TimescaleDB v2.30.2 on PostgreSQL 16.15. "
        "Database timezone was set to Africa/Kigali (+02). A pilot stream of 2,000 readings was validated with zero loss "
        "across generator, broker, subscriber, and database landing table. The full historical stream was captured "
        "and loaded via CSV bulk COPY in 216.22 seconds, achieving the exact target of 10,696,848 emitted rows."
    )

    doc.add_heading("3. Data Quality & Cleaning", level=1)
    doc.add_paragraph(
        "Before cleaning, the system identified 10,643,735 unique readings, 53,113 duplicate extra copies (0.4990%), "
        "213,090 delayed unique readings (2.0020%), and 108,265 missing readings (1.0069%), achieving 98.9931% completeness. "
        "Investigation confirmed 0 conflicting values among duplicate pairs. Repeatable deduplication purged all 53,113 extra "
        "copies, leaving exactly 10,643,735 verified unique rows."
    )

    doc.add_heading("4. Chunk Interval Benchmark (Section 3)", level=1)
    doc.add_paragraph(
        "Three hypertable configurations were benchmarked across 5 representative query workloads: 3-hour (225 chunks), "
        "1-day (29 chunks), and 1-week (5 chunks). While 3-hour chunks excelled at short 1-day time slices (Q1: 70.57 ms), "
        "the 1-day chunk interval demonstrated superior balance across full-period meter aggregations (Q3: 1,125.81 ms vs 3,903.10 ms) "
        "and single-meter index scans (Q4: 6.56 ms), proving to be the optimal operational configuration."
    )

    doc.add_heading("5. Compression Strategies (Section 4)", level=1)
    doc.add_paragraph(
        "Two compression configurations were evaluated against the 1-day baseline (1,435.52 MB): unsegmented (438.96 MB, 69.42% savings) "
        "and segmented by meter_id (501.88 MB, 65.04% savings). Segmented compression achieved outstanding query acceleration: "
        "Q3 improved by +79.46% (231.29 ms), Q4 single-meter lookup ran in 6.16 ms (+6.10%), and full scans ran in 348.40 ms (+60.86%). "
        "Consequently, 1-day chunks with meter segmentation is selected as the recommended storage architecture."
    )

    doc.add_heading("6. Utility Data Investigation & Visualizations (Section 5)", level=1)
    doc.add_paragraph(
        "Total network consumption expanded steadily by ~0.3% per day, reaching a maximum 15-minute peak of 7,993.69 kW on 28 September at 13:00 CAT. "
        "Region 8 ranked #1 in consumption (503,475.83 kWh), heavily driven by industrial load (67.24%). "
        "Evaluated against 40 injected abnormal events, our baseline deviation rule detected 24 events with 0 false flags (0.00% false flag rate)."
    )

    if os.path.exists("demand_patterns.png"):
        doc.add_paragraph("Figure 1: Network Demand Trend and Weekday vs Weekend Hourly Load Profiles")
        doc.add_picture("demand_patterns.png", width=Inches(6.2))

    doc.add_heading("7. Continuous Aggregates & Freshness (Section 6 & 7)", level=1)
    doc.add_paragraph(
        "Continuous aggregates were created for 15-minute regional, hourly regional, and daily meter buckets. "
        "Across 10 measured refreshes, the 5-panel operational dashboard saw median render time drop from 1,445.17 ms to 40.00 ms (a 36.1x speedup) "
        "with exactly 0.000000 kWh divergence across all panels. The freshness timeline test confirmed proper visibility "
        "of raw and materialized state across all stages."
    )

    doc.add_heading("8. Day-Ahead Regional Forecasting (Section 8)", level=1)
    doc.add_paragraph(
        "A Ridge Regression model trained on 14 days and tested across 1,680 test observations (22–28 September) achieved a test MAE of 4.52 kWh "
        "and RMSE of 5.74 kWh, delivering an 89.43% error reduction compared to the naive persistence baseline (42.71 kWh MAE). "
        "Peak-hour prediction error in high-demand industrial sectors dropped from >250 kWh to <8 kWh."
    )

    if os.path.exists("forecast_plot.png"):
        doc.add_paragraph("Figure 2: Actual vs Baseline vs Model Forecast (Region 8, 7 Test Days)")
        doc.add_picture("forecast_plot.png", width=Inches(6.2))

    doc.add_heading("9. Executive Management Recommendation (Section 5.5)", level=1)
    doc.add_paragraph(
        "1. Adopt TimescaleDB 1-day chunk interval with segmented compression (meter_id) to cut storage by 65.04% and accelerate meter queries to 6.16 ms.\n"
        "2. Deploy continuous aggregates with 30-minute refresh intervals covering trailing 7 days, achieving 36.1x dashboard speedup while seamlessly capturing delayed telemetry.\n"
        "3. Implement the Ridge Regression model for grid dispatch planning, eliminating weekend-to-weekday peak transition errors and keeping peak load discrepancies below 8 kWh."
    )

    doc.save("Assignment_I_SEGA_Final_Report.docx")
    print("Report saved successfully to 'Assignment_I_SEGA_Final_Report.docx'.")

if __name__ == "__main__":
    main()
